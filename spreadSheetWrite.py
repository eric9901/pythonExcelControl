from openpyxl import Workbook #for create and write spreadsheet
from openpyxl.worksheet.table import Table, TableStyleInfo #for create table
def createXlsxFile():
 wb = Workbook() # 創建一個空白活頁簿物件
 return wb

def appendValue(invalue):
     ws.append(invalue)  # 向下新增一列並連續插入值 可以直接塞value e.g("asd","sd") or list variable

def columeInValue(pos,value):
 ws[pos]=value;

def tableCreate(tableName,pos): #pos=位置 e.g.A1:B5
    tab = Table(displayName="tableName", ref=pos)  # ref is the table range
    style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,  # the style of table
                           showLastColumn=False, showRowStripes=True, showColumnStripes=True)
    tab.tableStyleInfo = style  # apply the style to table
    ws.add_table(tab)  # add table to word sheet

wb=createXlsxFile()#craete workbook
ws = wb.active# 選取正在工作中的表單
# 指定值給 A1 儲存格
#columeInValue('A1','我是儲存格')

frust=(["frust", "2014","2015","2016","2017"])  #Colume heading must be string, empty is also not allowed
appendValue(frust)
#table creating
data= [["apples",1000,5000,8000,600],["orange",1000,300,40,6000]]
for row in data:
    appendValue(row)
tableCreate("frust","A1:E2")
wb.save('create_sample.xlsx') #create an xlsx file
# teast case1 run 2 time with same file name
#only 1 file created
# input some data between two run, the file is recreate to its origion form.

